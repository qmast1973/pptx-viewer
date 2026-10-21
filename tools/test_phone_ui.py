"""폰 화면 시험 (헤드리스 Chrome 393×852, 소리 끔, 실제 입력 경로 CDP Input.dispatchTouchEvent).
2026-10-21 사용자 요청 확인용: 폰에서도 썸네일(슬라이드 아래), 전체화면에서는 숨김, 글자를 칠 때마다 바로 찾기.
  python tools/test_phone_ui.py [폴더]      (큰 덱 시험용 .pptx 는 python-pptx 가 있으면 임시로 만든다)
기대: 보통 화면 썸네일 줄 104px(4장·20장), 150장은 세로 목록 200px / 전체화면 썸네일 없음 / 한 글자 칠 때마다 결과가 바뀌고 입력칸 초점 유지 /
      20장 덱 썸네일 줄을 가로로 밀면 줄만 움직이고 쪽은 그대로 / 썸네일 누르면 그 쪽으로 이동."""
import asyncio, json, urllib.request, subprocess, tempfile, sys, shutil, os
from pathlib import Path
import websockets

ROOT = sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).resolve().parent.parent)
PORT, DBG = 8795, 9335
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'
fails = []


def check(name, ok, info=""):
    print(("  OK  " if ok else " FAIL ") + name, info)
    if not ok:
        fails.append(name)


def make_deck(path, n):
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
    for i in range(n):
        s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = f"슬라이드 {i + 1}"
        tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)); tb.text_frame.text = ("매출 보고 " if i % 7 == 0 else "일반 ") + str(i + 1)
    p.save(path)


async def main():
    work = tempfile.mkdtemp()
    decks = {}
    try:
        for n in (20, 150):
            decks[n] = os.path.join(work, f"d{n}.pptx"); make_deck(decks[n], n)
    except ImportError:
        print("python-pptx 없음 — 20장·150장 덱 시험은 건너뜀")
    srv = subprocess.Popen([sys.executable, "-m", "http.server", str(PORT), "--directory", ROOT], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    prof = tempfile.mkdtemp()
    proc = subprocess.Popen([CHROME, '--headless=new', f'--remote-debugging-port={DBG}', f'--user-data-dir={prof}', '--disable-cache', '--mute-audio', '--no-first-run', f'http://127.0.0.1:{PORT}/'])
    try:
        await asyncio.sleep(3)
        tabs = json.loads(urllib.request.urlopen(f'http://127.0.0.1:{DBG}/json').read().decode())
        target = [t for t in tabs if t.get('type') == 'page' and str(PORT) in t.get('url', '')][0]
        async with websockets.connect(target['webSocketDebuggerUrl'], max_size=None) as ws:
            n = 0

            async def send(m, p=None):
                nonlocal n
                n += 1
                await ws.send(json.dumps({'id': n, 'method': m, 'params': p or {}}))
                while True:
                    r = json.loads(await ws.recv())
                    if r.get('id') == n:
                        return r.get('result', {})

            async def ev(js):
                return (await send('Runtime.evaluate', {'expression': js, 'returnByValue': True, 'awaitPromise': True})).get('result', {}).get('value')

            async def touch(kind, pts):
                await send('Input.dispatchTouchEvent', {'type': kind, 'touchPoints': [{'x': x, 'y': y, 'id': i} for i, (x, y) in enumerate(pts)]})

            await send('Emulation.setDeviceMetricsOverride', {'width': 393, 'height': 852, 'deviceScaleFactor': 2, 'mobile': True})
            await send('Emulation.setTouchEmulationEnabled', {'enabled': True, 'maxTouchPoints': 5})

            async def wait_count(c):
                for _ in range(120):
                    if await ev(f"Boolean(window.pptxViewer) && window.pptxViewer.state.count=={c}"):
                        break
                    await asyncio.sleep(0.5)
                await asyncio.sleep(1.5)

            async def load(path, c):
                doc = await send('DOM.getDocument')
                node = await send('DOM.querySelector', {'nodeId': doc['root']['nodeId'], 'selector': '#fileIn'})
                await send('DOM.setFileInputFiles', {'files': [path], 'nodeId': node['nodeId']})
                await wait_count(c)

            def geo():
                return ev("(()=>{const t=document.querySelector('.pptxv-thumbs');const r=t.getBoundingClientRect();"
                          "return {disp:getComputedStyle(t).display,h:Math.round(r.height),w:Math.round(r.width)}})()")

            await wait_count(4)
            g = await geo()
            check("폰 보통 화면: 썸네일 줄이 보이고 가로 폭 가득, 높이 104", g['disp'] == 'flex' and g['h'] == 104 and g['w'] == 393, g)
            await ev("document.querySelectorAll('.pptxv-thumb')[2].click()"); await asyncio.sleep(0.4)
            check("썸네일 누르면 그 쪽으로 이동", await ev("window.pptxViewer.state.slide") == 2)
            await ev("window.pptxViewer.go(0)"); await asyncio.sleep(0.3)
            await ev("document.getElementById('q').focus()")
            res = "({title:document.getElementById('resTitle').textContent,focus:document.activeElement.id})"
            await send('Input.insertText', {'text': '매'}); await asyncio.sleep(0.7); a = await ev(res)
            await send('Input.insertText', {'text': '출'}); await asyncio.sleep(0.7); b = await ev(res)
            check("한 글자 칠 때마다 찾기(찾기 버튼 없이) + 입력칸 초점 유지", a['title'] != b['title'] and '매출' in b['title'] and b['focus'] == 'q', (a['title'], b['title']))
            await ev("window.pptxViewer.toggleFull()"); await asyncio.sleep(1)
            g = await geo()
            check("전체화면에서는 썸네일 숨김", g['disp'] == 'none', g)
            await ev("window.pptxViewer.toggleFull()"); await asyncio.sleep(1)
            if 20 in decks:
                await load(decks[20], 20)
                g = await geo()
                check("20장: 가로 줄 104", g['disp'] == 'flex' and g['h'] == 104, g)
                await ev("window.scrollTo(0,0)"); await asyncio.sleep(0.3)
                tr = await ev("(()=>{const r=document.querySelector('.pptxv-thumbs').getBoundingClientRect();return [r.left+r.width/2,r.top+r.height/2]})()")
                s0 = await ev("window.pptxViewer.state.slide")
                await touch('touchStart', [(tr[0] + 120, tr[1])])
                for i in range(1, 11):
                    await touch('touchMove', [(tr[0] + 120 - 24 * i, tr[1])]); await asyncio.sleep(0.016)
                await touch('touchEnd', []); await asyncio.sleep(0.6)
                sc = await ev("document.querySelector('.pptxv-thumbs-list').scrollLeft")
                check("20장: 썸네일 줄을 가로로 밀면 줄만 움직이고 쪽은 그대로", sc > 50 and await ev("window.pptxViewer.state.slide") == s0, sc)
            if 150 in decks:
                await load(decks[150], 150)
                g = await geo()
                check("150장: 세로 목록 200", g['disp'] == 'flex' and g['h'] == 200, g)
                st = await ev("(()=>{const s=document.getElementById('stage').getBoundingClientRect(),v=document.querySelector('.pptxv-viewport').getBoundingClientRect();return [Math.round(s.height),Math.round(v.height)]})()")
                check("150장: 슬라이드 칸 = 슬라이드 비율 높이 + 200", st[0] - st[1] == 200, st)
                await ev("document.querySelectorAll('.pptxv-thumb')[5].click()"); await asyncio.sleep(0.4)
                check("150장: 썸네일 누르면 그 쪽으로 이동", await ev("window.pptxViewer.state.slide") == 5)
    finally:
        proc.terminate(); srv.terminate(); await asyncio.sleep(1)
        shutil.rmtree(prof, ignore_errors=True); shutil.rmtree(work, ignore_errors=True)
    print("\n>> 모두 통과 <<" if not fails else f"\n>> 실패 {len(fails)}건: {fails}")
    sys.exit(1 if fails else 0)

asyncio.run(main())
