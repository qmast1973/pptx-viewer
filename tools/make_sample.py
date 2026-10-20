"""화면 첫 상태에 보여 줄 예시 발표자료(sample/예시_발표.pptx)를 만든다. python-pptx 필요."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parent.parent / "sample" / "예시_발표.pptx"
SLIDES = [
    ("분기 실적 보고", ["2026년 3분기 · 영업팀", "매출과 고객 수 변화를 정리했습니다."]),
    ("매출 요약", ["3분기 매출 42억 원, 전 분기보다 12% 증가", "신규 고객 매출 비중 31%", "온라인 주문이 처음으로 매장 주문을 넘음"]),
    ("고객 현황", ["신규 고객 1,240명", "재구매율 58% (목표 55%)", "고객 문의 응답 시간 평균 3시간"]),
    ("다음 분기 계획", ["온라인 주문 화면 개선", "재구매 고객 혜택 확대", "매출 목표 46억 원"]),
]


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]
    for title, lines in SLIDES:
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Emu(700000), Emu(500000), Emu(10800000), Emu(1100000))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text, r.font.size, r.font.bold = title, Pt(40), True
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        for i, line in enumerate(lines):
            b = s.shapes.add_textbox(Emu(900000), Emu(2000000 + i * 1000000), Emu(10400000), Emu(800000))
            rr = b.text_frame.paragraphs[0].add_run()
            rr.text, rr.font.size = line, Pt(24)
            rr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT, OUT.stat().st_size)


if __name__ == "__main__":
    main()
